"""Builds Agent-Registry-Test-Cases-v1.xlsx (Summary, Test Cases, Defect Log, Fields) from the three CSV files in the repo root.
(Same layout as the test-maker workbook.)

    python3 tools/test_cases/build_xlsx.py     # needs openpyxl (pip install openpyxl)
"""
import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

R = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")) + os.sep
read = lambda f: list(csv.reader(open(R + f, encoding="utf-8-sig")))
cases, fields, defects = read("Agent-Registry-Test-Cases-v1.csv"), read("Agent-Registry-Test-Case-Fields-v1.csv"), read("Agent-Registry-Defect-Log-v1.csv")

F = Font(name="Arial", size=10)
H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="4F46E5")
EDIT = PatternFill("solid", fgColor="FFFFCC")
thin = Side(style="thin", color="D9D9D9")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

def fill(ws, rows, widths, edit_cols=()):
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.alignment, cell.border = (H if r == 1 else F), wrap, B
            if r == 1:
                cell.fill = HF
            elif c in edit_cols:
                cell.fill = EDIT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

wb = Workbook()
# ── Summary ──
s = wb.active
s.title = "Summary"
n = 1000  # room for cases added later
rng = lambda col: f"'Test Cases'!${col}$2:${col}${n}"
s["A1"] = "Agent Registry: test cases v1"; s["A1"].font = Font(name="Arial", size=14, bold=True)
s["A2"] = ("How to use: on the 'Test Cases' sheet fill only the yellow columns (Manual Status, Actual Result, Failure Description, "
           "Defect ID, Tested By, Test Date, Build / Version). Log each failure on the 'Defect Log' sheet and put its id in Defect ID. "
           "'Fields' explains every column. Numbers below update by themselves.")
s["A2"].alignment = wrap; s["A2"].font = F; s.merge_cells("A2:H2"); s.row_dimensions[2].height = 48
s["A4"], s["B4"] = "Manual status", "Cases"
for i, st in enumerate(["Not Run", "Pass", "Fail", "Blocked", "N/A"], 5):
    s[f"A{i}"] = st; s[f"B{i}"] = f'=COUNTIF({rng("M")},A{i})'
s["A10"], s["B10"] = "Total", f'=COUNTA({rng("A")})'
s["A11"], s["B11"] = "Executed (Pass + Fail)", "=B6+B7"
s["A12"], s["B12"] = "Pass rate of executed", '=IF(B11=0,"-",B6/B11)'; s["B12"].number_format = "0.0%"
s["D4"], s["E4"], s["F4"], s["G4"] = "Priority", "Cases", "Passed", "Failed"
for i, p in enumerate(["P1", "P2", "P3"], 5):
    s[f"D{i}"] = p; s[f"E{i}"] = f'=COUNTIF({rng("E")},D{i})'
    s[f"F{i}"] = f'=COUNTIFS({rng("E")},D{i},{rng("M")},"Pass")'; s[f"G{i}"] = f'=COUNTIFS({rng("E")},D{i},{rng("M")},"Fail")'
mods = sorted({r[1] for r in cases[1:]}, key=[r[1] for r in cases[1:]].index)
s["A15"], s["B15"], s["C15"], s["D15"], s["E15"] = "Module", "Cases", "Passed", "Failed", "Not run"
for i, m in enumerate(mods, 16):
    s[f"A{i}"] = m; s[f"B{i}"] = f'=COUNTIF({rng("B")},A{i})'
    for col, st in (("C", "Pass"), ("D", "Fail"), ("E", "Not Run")):
        s[f"{col}{i}"] = f'=COUNTIFS({rng("B")},A{i},{rng("M")},"{st}")'
for row in s.iter_rows(min_row=4):
    for cell in row:
        if cell.value is not None:
            cell.font = F; cell.border = B
for ref in ("A4", "B4", "D4", "E4", "F4", "G4", "A15", "B15", "C15", "D15", "E15"):
    s[ref].font, s[ref].fill = H, HF
s.column_dimensions["A"].width = 44
for c in "BCDEFGH":
    s.column_dimensions[c].width = 14

# ── Test Cases ──
t = wb.create_sheet("Test Cases")
fill(t, cases, [11, 26, 30, 14, 8, 12, 34, 60, 60, 13, 38, 12, 13, 34, 40, 11, 14, 12, 18, 9, 30, 60], edit_cols=range(13, 20))
last = len(cases) + 200
for col, opts in (("M", "Not Run,Pass,Fail,Blocked,N/A"), ("E", "P1,P2,P3"), ("F", "Functional,Negative,Security,Integration,UI")):
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True); dv.error = "Pick a value from the list"
    t.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
for val, color in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C")):
    t.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
t.cell(1, 18).comment = None
for r in range(2, len(cases) + 1):
    t.cell(r, 18).number_format = "yyyy-mm-dd"

# ── Defect Log ── (one example row shows the expected format)
d = wb.create_sheet("Defect Log")
example = ["DEF-000", "EXAMPLE (delete this row): Finish stays disabled after every step is done", "RET-001", "Medium",
           "1. Governance tab, Retire this agent. 2. Complete every step.", "Finish becomes enabled", "Finish stays disabled",
           "", "Open", "Developer name", "main @ abc1234", "", "", "2026-10-12", "", "Example only"]
fill(d, defects + ([example] if len(defects) == 1 else []), [10, 40, 11, 11, 50, 34, 34, 22, 12, 16, 22, 22, 13, 12, 12, 30])
for col, opts in (("D", "Critical,High,Medium,Low"), ("I", "Open,In Progress,Fixed,Retest,Closed,Won't Fix"), ("M", "Pass,Fail")):
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True); d.add_data_validation(dv); dv.add(f"{col}2:{col}500")

# ── Fields ──
fill(wb.create_sheet("Fields"), fields, [22, 80, 60, 20])

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(R + "Agent-Registry-Test-Cases-v1.xlsx")
print("rows:", len(cases) - 1)
